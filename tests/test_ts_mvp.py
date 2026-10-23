"""Network-free vertical tests: synthetic API record + generated fictional protocol."""

import csv
import io
import json
from pathlib import Path
from urllib.error import HTTPError, URLError

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from test_extraction import pdf_bytes

from trial_design_workflow import inputs
from trial_design_workflow.excel import TS_COLUMNS, workbook_bytes
from trial_design_workflow.inputs import ASSETS, Fact, collect_inputs, reconcile, registry_facts
from trial_design_workflow.ts_cli import main, run
from trial_design_workflow.ts_registry import build_candidates, ct_lookup, load_ct, load_registry

FIXTURE = Path(__file__).parent / "fixtures/registry-synthetic-v1.json"


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def blocked(*args, **kwargs):
        raise AssertionError("Automated fixture tests must not access the network")

    monkeypatch.setattr(inputs, "urlopen", blocked)


@pytest.fixture
def payload():
    return FIXTURE.read_bytes()


@pytest.fixture
def pdf(tmp_path):
    content = pdf_bytes(
        [
            "SYNTHETIC PROTOCOL ONLY\nNCT00000001\nSponsor trial ID: SYN-101\n",
            "Page: 1 of 1\nProtocol title:\nFictional study\nSubstance names: fictional product\n"
            "Trial ID: SYN-101 Version: 1.0\nTrial phase: 2\n"
            "Approximately 240 subjects will be screened to achieve "
            "180 subjects randomly assigned.\n",
        ]
    )
    path = tmp_path / "synthetic.pdf"
    path.write_bytes(content)
    return path


def observations(value, source="protocol", key="protocol_title"):
    return Fact(
        field_key=key,
        proposed_value=value,
        source_type=source,
        source_locator="PDF page=1" if source == "protocol" else "protocolSection.test",
        evidence_text=value if source == "protocol" else "",
        source_id="synthetic",
        extraction_status="proposed" if source == "protocol" else "parsed",
    )


def test_reconcile_agreement_single_source_and_conflict():
    facts = (observations("Fictional study"), observations("Fictional Study", "clinicaltrials_gov"))
    agreed = reconcile(facts)[0]
    assert agreed.result == "agreement"
    assert len(agreed.values) == 1
    assert len(agreed.values[0].observations) == 2
    assert {f.proposed_value for f in agreed.values[0].observations} == {
        "Fictional study",
        "Fictional Study",
    }
    assert reconcile(facts[:1])[0].result == "single_source"
    conflict = reconcile((*facts, observations("Different title")))[0]
    assert conflict.result == "conflict"
    assert len(conflict.values) == 2
    assert reconcile(tuple(reversed(facts))) == reconcile(facts)


@pytest.mark.parametrize("mode", ["registry", "protocol", "hybrid"])
def test_modes_share_facts_and_one_excel_path(mode, pdf, tmp_path):
    kwargs = {"nct_id": "NCT00000001", "registry_json": FIXTURE} if mode != "protocol" else {}
    if mode != "registry":
        kwargs["protocol"] = pdf
    study = collect_inputs(**kwargs)
    assert study.input_type == mode
    assert study.facts and all(f.review_status == "needs_review" for f in study.facts)
    destination = run(output=tmp_path / mode, **kwargs)
    assert list(destination.parent.iterdir()) == [destination]
    book = load_workbook(destination)
    assert book.sheetnames == ["TS_Draft", "Parameter_Review", "Evidence_QC"]
    assert tuple(c.value for c in book["TS_Draft"][1]) == TS_COLUMNS
    assert book["TS_Draft"].max_row == book["Parameter_Review"].max_row == 130
    rows = {r[3]: r for r in book["TS_Draft"].iter_rows(min_row=2, values_only=True)}
    assert rows["TITLE"][5].casefold() == "fictional study"
    assert rows["ACTSUB"][5] == ("120" if mode != "protocol" else None)
    assert rows["PLANSUB"][5] == ("180" if mode != "registry" else None)
    assert rows["ADAPT"][5] is None
    assert all(c.value == "needs_review" for c in book["Parameter_Review"]["J"][1:])
    if mode == "hybrid":
        review = {r[0]: r for r in book["Parameter_Review"].iter_rows(min_row=2, values_only=True)}
        assert review["TITLE"][6] == "agreement"
        assert review["TPHASE"][2] == "conflict"  # raw 2 versus PHASE2 is deliberately not hidden
        assert rows["TPHASE"][5] is None
    book.close()


def test_registry_live_adapter_can_be_stubbed_without_network(monkeypatch, payload):
    monkeypatch.setattr(inputs, "urlopen", lambda *a, **kw: io.BytesIO(payload))
    study = collect_inputs(nct_id="NCT00000001")
    assert study.fetched_at.tzinfo is not None
    assert len(study.payload_sha256) == 64
    assert study.input_type == "registry"


def test_conflicting_enrollment_is_retained_with_both_sources(pdf, payload, tmp_path):
    snapshot = json.loads(payload)
    snapshot["protocolSection"]["designModule"]["enrollmentInfo"] = {
        "count": 200,
        "type": "ESTIMATED",
    }
    path = tmp_path / "synthetic-snapshot.json"
    path.write_text(json.dumps(snapshot))
    study = collect_inputs(nct_id="NCT00000001", protocol=pdf, registry_json=path)
    field = next(f for f in study.facts if f.field_key == "planned_enrollment")
    assert field.result == "conflict"
    assert {v.proposed_value for v in field.values} == {"180", "200"}
    registry, ct = (
        load_registry(ASSETS / "ts-parameters-v1.csv"),
        load_ct(ASSETS / "ct-2025-03-28.csv"),
    )
    candidates = build_candidates(study, registry, ct)
    planned = next(c for c in candidates if c.parameter.TSPARMCD == "PLANSUB")
    assert planned.recommendation == "conflict" and planned.draft_value is None


def test_identity_mismatch_fails_before_workbook_creation(pdf, payload, tmp_path):
    with pytest.raises(ValueError, match="response NCT ID"):
        registry_facts(payload, "NCT00000002")
    changed = json.loads(payload)
    changed["protocolSection"]["identificationModule"]["nctId"] = "NCT00000002"
    path = tmp_path / "other.json"
    path.write_text(json.dumps(changed))
    with pytest.raises(ValueError, match="Protocol NCT identity conflicts"):
        run(nct_id="NCT00000002", registry_json=path, protocol=pdf, output=tmp_path / "failed")
    assert not (tmp_path / "failed").exists()


@pytest.mark.parametrize(
    "mutation",
    [
        lambda p: p["protocolSection"]["designModule"]["enrollmentInfo"].update(
            count="not numeric"
        ),
        lambda p: p["protocolSection"]["designModule"]["enrollmentInfo"].update(type="UNKNOWN"),
        lambda p: p["protocolSection"]["identificationModule"].update(nctId="invalid"),
    ],
)
def test_bad_api_payloads_rejected(payload, mutation):
    document = json.loads(payload)
    mutation(document)
    with pytest.raises(ValidationError):
        registry_facts(json.dumps(document).encode(), "NCT00000001")


def test_missing_api_fields_are_not_fabricated():
    facts = registry_facts(
        b'{"protocolSection":{"identificationModule":{"nctId":"NCT00000001"}}}', "NCT00000001"
    )
    assert [f.field_key for f in facts] == ["nct_id"]


@pytest.mark.parametrize(
    "exception", [URLError("offline"), HTTPError("test", 404, "missing", {}, None)]
)
def test_network_failures_are_actionable(monkeypatch, exception):
    def fail(*args, **kwargs):
        raise exception

    monkeypatch.setattr(inputs, "urlopen", fail)
    with pytest.raises(ValueError, match="ClinicalTrials.gov"):
        collect_inputs(nct_id="NCT00000001")


def test_ct_exact_lookup_and_missing_codelists():
    ct = load_ct(ASSETS / "ct-2025-03-28.csv")
    assert ct_lookup(ct, "C66737", "PHASE2")[0] == "PHASE II TRIAL"
    assert ct_lookup(ct, "C66737", "unknown")[0] is None
    assert ct_lookup(ct, "NOT_LOADED", "unknown")[0] is None
    assert ct_lookup(ct, "UNVERIFIED", "anything")[0] is None
    # QUADRUPLE must not be guessed as DOUBLE BLIND when the snapshot has no such term.
    assert ct_lookup(ct, "C66735", "QUADRUPLE")[0] is None


def test_registry_duplicates_versions_and_release_mismatch(tmp_path):
    path = ASSETS / "ts-parameters-v1.csv"
    with path.open(newline="") as stream:
        reader = csv.DictReader(stream)
        fields, rows = reader.fieldnames, list(reader)
    duplicate = tmp_path / "duplicate.csv"
    with duplicate.open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows([*rows, rows[0]])
    with pytest.raises(ValueError, match="Duplicate TSPARMCD"):
        load_registry(duplicate)
    registry = load_registry(path)
    ct = load_ct(ASSETS / "ct-2025-03-28.csv")
    with pytest.raises(ValueError, match="CT release"):
        build_candidates(
            collect_inputs(nct_id="NCT00000001", registry_json=FIXTURE),
            registry,
            ct.model_copy(update={"release": "other"}),
        )


def test_applicability_exclusion_external_and_unsupported_are_visible():
    registry = load_registry(ASSETS / "ts-parameters-v1.csv")
    title = next(e for e in registry.entries if e.TSPARMCD == "TITLE")
    study = collect_inputs(nct_id="NCT00000001", registry_json=FIXTURE)
    ct = load_ct(ASSETS / "ct-2025-03-28.csv")
    for updates, expected in (
        ({"applicability_rule": "exclude"}, "not_applicable"),
        ({"applicability_rule": "external"}, "needs_external_source"),
        ({"automation_status": "unsupported"}, "unsupported"),
        ({"TSPARM": "Wrong label"}, "needs_review"),
    ):
        changed = registry.model_copy(update={"entries": (title.model_copy(update=updates),)})
        candidates = build_candidates(study, changed, ct)
        assert len(candidates) == 1
        assert candidates[0].recommendation == expected
        assert candidates[0].draft_value is None


def test_formula_injection_is_literal_and_long_text_blocks(payload, tmp_path):
    document = json.loads(payload)
    document["protocolSection"]["identificationModule"]["officialTitle"] = "=1+1"
    path = tmp_path / "formula.json"
    path.write_text(json.dumps(document))
    output = run(nct_id="NCT00000001", registry_json=path, output=tmp_path / "formula.xlsx")
    book = load_workbook(output)
    title = next(row for row in book["TS_Draft"].iter_rows(min_row=2) if row[3].value == "TITLE")
    assert title[5].value == "=1+1" and title[5].data_type == "s"
    book.close()
    document["protocolSection"]["identificationModule"]["officialTitle"] = "x" * 201
    path.write_text(json.dumps(document))
    study = collect_inputs(nct_id="NCT00000001", registry_json=path)
    registry, ct = (
        load_registry(ASSETS / "ts-parameters-v1.csv"),
        load_ct(ASSETS / "ct-2025-03-28.csv"),
    )
    candidate = next(
        c for c in build_candidates(study, registry, ct) if c.parameter.TSPARMCD == "TITLE"
    )
    assert candidate.draft_value is None and "project review threshold" in candidate.message


def test_repeatability_all_evidence_and_no_overwrite(tmp_path, pdf):
    first = run(nct_id="NCT00000001", registry_json=FIXTURE, protocol=pdf, output=tmp_path / "one")
    second = run(nct_id="NCT00000001", registry_json=FIXTURE, protocol=pdf, output=tmp_path / "two")
    a, b = load_workbook(first), load_workbook(second)
    for sheet in a.sheetnames:
        assert list(a[sheet].values) == list(b[sheet].values)
    assert all(cell.value is not None for cell in a["Parameter_Review"]["M"][1:])
    messages = [r[11] for r in a["Evidence_QC"].iter_rows(min_row=2, values_only=True)]
    assert any("STANDARDS_GAP" in str(message) for message in messages)
    assert any(r[2] == "payload_sha256" for r in a["Evidence_QC"].iter_rows(values_only=True))
    a.close()
    b.close()
    with pytest.raises(ValueError, match="already exists"):
        run(nct_id="NCT00000001", registry_json=FIXTURE, output=first)


def test_cli_usage_error_and_no_sensitive_echo(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("sys.argv", ["trial-ts", "--output", str(tmp_path / "unused")])
    assert main() == 2
    assert json.loads(capsys.readouterr().err)["code"] == "TS_DRAFT_ERROR"
    assert not (tmp_path / "unused").exists()


def test_excel_rejects_unrepresentable_source_instead_of_silent_truncation(payload):
    study = collect_inputs(nct_id="NCT00000001", registry_json=FIXTURE)
    field = reconcile((observations("x" * 32768),))[0]
    changed = study.model_copy(update={"facts": (field,)})
    registry = load_registry(ASSETS / "ts-parameters-v1.csv")
    ct = load_ct(ASSETS / "ct-2025-03-28.csv")
    candidates = build_candidates(changed, registry, ct)
    with pytest.raises(ValueError, match="Unrepresentable Excel text"):
        workbook_bytes(changed, registry, ct, candidates)


@pytest.mark.parametrize(
    "date_value, expected",
    [
        ("2025", "2025"),
        ("2025-02", "2025-02"),
        ("2024-02-29", "2024-02-29"),
        ("2025-02-29", None),
        ("2025-19", None),
        ("0000", None),
    ],
)
def test_date_precision_is_preserved_and_invalid_calendar_dates_blocked(date_value, expected):
    study = collect_inputs(nct_id="NCT00000001", registry_json=FIXTURE)
    field = reconcile((observations(date_value, "clinicaltrials_gov", "study_start_date"),))[0]
    study = study.model_copy(update={"facts": (field,)})
    registry = load_registry(ASSETS / "ts-parameters-v1.csv")
    ct = load_ct(ASSETS / "ct-2025-03-28.csv")
    result = next(
        c for c in build_candidates(study, registry, ct) if c.parameter.TSPARMCD == "SSTDTC"
    )
    assert result.draft_value == expected


def test_single_source_with_internal_conflict_never_chooses_one_value():
    field = reconcile((observations("Alpha"), observations("Beta")))[0]
    assert field.result == "conflict"
    assert {v.proposed_value for v in field.values} == {"Alpha", "Beta"}


def test_all_candidate_codes_and_names_are_in_pinned_nci_snapshot():
    registry = load_registry(ASSETS / "ts-parameters-v1.csv")
    ct = load_ct(ASSETS / "ct-2025-03-28.csv")
    codes = {t.value: t.code for t in ct.terms if t.codelist == "C66738"}
    names = {t.code: t.value for t in ct.terms if t.codelist == "C67152"}
    assert {entry.TSPARMCD for entry in registry.entries} == set(codes)
    assert all(entry.TSPARM == names[codes[entry.TSPARMCD]] for entry in registry.entries)
    assert all(entry.applicability_rule == "review" for entry in registry.entries)


def test_protocol_mode_never_uses_registry_transport(pdf, monkeypatch):
    def fail(*args, **kwargs):
        raise AssertionError("Protocol-only mode must not call the API adapter")

    monkeypatch.setattr(inputs, "fetch_record", fail)
    assert collect_inputs(protocol=pdf).input_type == "protocol"
