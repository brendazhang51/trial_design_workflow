"""Read locally supplied CDISC tables; no download, approval or study interpretation."""

import re
from collections import defaultdict
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from pydantic import Field
from pypdf import PdfReader

from ..schemas import Model, Text


class ResponseLink(Model):
    codelist: Text
    codelist_name: Text
    term_code: Text
    value: Text
    locator: Text
    source_value_kind: Literal["text", "numeric"] = "text"


class CodetableParameter(Model):
    concept_code: Text
    tsparmcd: Text
    tsparm: Text
    locators: tuple[Text, ...]
    responses: tuple[ResponseLink, ...]


class CodetableSnapshot(Model):
    release: Text
    parameters: tuple[CodetableParameter, ...] = Field(min_length=1)


def read_codetable(path: Path, release: str) -> CodetableSnapshot:
    """Accept the explicit three-column-parameter/response-pair export layout only."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1 or release not in workbook.worksheets[0].title:
            raise ValueError("Codetable worksheet must identify the selected release")
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        if "C66738" not in str(headers[1]) or "C67152" not in str(headers[2]):
            raise ValueError("Expected TSPARMCD/TSPARM concept-paired columns")
        response_columns = {}
        for index in range(5, len(headers), 3):
            header = str(headers[index])
            match = re.search(r"codelist code\s*=\s*(C\d+)", header)
            if not match:
                raise ValueError("Unrecognized response-codelist header")
            name = " ".join(header.split("(", 1)[0].split())
            response_columns[index] = (match[1], name)
        parameters = {}
        for row_number, row in enumerate(rows, 2):
            if not any(v is not None for v in row):
                continue
            code, parameter, label = row[:3]
            if not all(
                isinstance(v, str) and v.strip() and not v.startswith("=")
                for v in (code, parameter, label)
            ):
                raise ValueError("Incomplete or formula-based parameter identity")
            if not re.fullmatch(r"C\d+", code):
                raise ValueError("Expected a parameter concept code")
            locator = f"{sheet.title}!A{row_number}:C{row_number}"
            entry = parameters.setdefault(
                parameter,
                {
                    "concept_code": code,
                    "tsparmcd": parameter,
                    "tsparm": label,
                    "locators": [],
                    "responses": [],
                },
            )
            if (entry["concept_code"], entry["tsparm"]) != (code, label):
                raise ValueError("Conflicting parameter identity in codetable")
            entry["locators"].append(locator)
            for index, (codelist, name) in response_columns.items():
                term_code, value = row[index - 1 : index + 1]
                if term_code is None and value is None:
                    continue
                if not isinstance(term_code, str) or not re.fullmatch(r"C(?:\d+|NEW)", term_code):
                    raise ValueError("Incomplete response code/value pair")
                value_kind = "numeric" if type(value) in {int, float} else "text"
                if value_kind == "numeric":
                    value = str(value)
                if not isinstance(value, str) or not value.strip() or value.startswith("="):
                    raise ValueError("Incomplete or formula-based response value")
                entry["responses"].append(
                    ResponseLink(
                        codelist=codelist,
                        codelist_name=name,
                        term_code=term_code,
                        value=value,
                        locator=f"{sheet.title}!row {row_number}, columns {index}:{index + 1}",
                        source_value_kind=value_kind,
                    )
                )
        return CodetableSnapshot(
            release=release,
            parameters=tuple(
                CodetableParameter.model_validate(parameters[key]) for key in sorted(parameters)
            ),
        )
    finally:
        workbook.close()


class ExtractedTerm(Model):
    codelist: Text
    code: Text
    value: Text
    definition: str
    pages: tuple[int, ...]


def ct_page_records(
    fragments: list[tuple[float, float, str]],
    codelist: str,
    page_number: int,
    previous: dict | None = None,
) -> tuple[list[dict], dict | None]:
    """Parse the pinned five-column PDF layout; keep cross-page continuation rows.

    Coordinates are specific to this source layout, not a general PDF-table parser.
    Source hashes and paired codetable checks guard against silently accepting drift.
    """
    headers = [y for x, y, text in fragments if text.strip() == "NCI Code" and x < 150]
    if len(headers) != 1:
        raise ValueError("CT table must have one recognizable column header")
    lines = defaultdict(list)
    for x, y, text in fragments:
        if 40 < y < headers[0] - 1:
            lines[round(y, 1)].append((x, text.strip()))
    records = []
    current = previous
    for _, cells in sorted(lines.items(), reverse=True):
        cells.sort()
        code = next((text for x, text in cells if x < 155 and re.fullmatch(r"C\d+", text)), None)
        if code:
            current = {
                "codelist": codelist,
                "code": code,
                "value_parts": [],
                "definition_parts": [],
                "pages": [],
            }
            records.append(current)
        if current is None:
            if any(x >= 155 for x, _ in cells):
                raise ValueError("Orphan CT row continuation")
            continue
        if page_number not in current["pages"]:
            current["pages"].append(page_number)
        current["value_parts"].extend(text for x, text in cells if 155 <= x < 270)
        current["definition_parts"].extend(text for x, text in cells if 380 <= x < 688)
    return records, current


def read_ct_pdf(path: Path, release: str, codelists: set[str]) -> tuple[ExtractedTerm, ...]:
    reader = PdfReader(path)
    if f"CDISC SDTM Controlled Terminology, {release}" not in reader.pages[0].extract_text():
        raise ValueError("CT PDF release differs from selected release")
    records = []
    current = None
    current_codelist = None
    for number, page in enumerate(reader.pages, 1):
        fragments = []

        def visit(text, cm, tm, font, size, collected=fragments):
            if text.strip():
                collected.append((cm[4] + tm[4], cm[5] + tm[5], text))

        text = page.extract_text(visitor_text=visit) or ""
        header = re.search(r"NCI Code: (C\d+), Codelist extensible:", text[:500])
        continuation = re.match(r"(C\d+)\s+\S+\s*\n", text)
        identity = header[1] if header else continuation[1] if continuation else None
        if identity not in codelists:
            current, current_codelist = None, None
            continue
        if identity != current_codelist:
            current = None
        parsed, current = ct_page_records(fragments, identity, number, current)
        records.extend(parsed)
        current_codelist = identity
    terms = tuple(
        ExtractedTerm(
            codelist=r["codelist"],
            code=r["code"],
            value=" ".join(r["value_parts"]),
            definition=" ".join(r["definition_parts"]),
            pages=tuple(r["pages"]),
        )
        for r in records
    )
    if {t.codelist for t in terms} != codelists:
        raise ValueError("Requested response/identity codelist is missing from CT PDF")
    if len({(t.codelist, t.code) for t in terms}) != len(terms):
        raise ValueError("Duplicate CT identity")
    return tuple(sorted(terms, key=lambda t: (t.codelist, t.code)))


class SourceDiscrepancy(Model):
    tsparmcd: Text
    locator: Text
    code: Text
    message: Text


def verify_codetable(
    snapshot: CodetableSnapshot, terms: tuple[ExtractedTerm, ...]
) -> tuple[SourceDiscrepancy, ...]:
    lookup = {(t.codelist, t.code): t.value for t in terms}
    discrepancies = []
    for parameter in snapshot.parameters:
        for codelist, code, value in (
            ("C66738", parameter.concept_code, parameter.tsparmcd),
            ("C67152", parameter.concept_code, parameter.tsparm),
        ):
            if lookup.get((codelist, code)) != value:
                discrepancies.append(
                    SourceDiscrepancy(
                        tsparmcd=parameter.tsparmcd,
                        locator=parameter.locators[0],
                        code="CODETABLE_IDENTITY_MISMATCH",
                        message=f"Codetable identity {codelist}/{code}/{value} "
                        "differs from selected CT; "
                        "retain CT identity and block this association pending source review.",
                    )
                )
        for response in parameter.responses:
            if lookup.get((response.codelist, response.term_code)) != response.value:
                discrepancies.append(
                    SourceDiscrepancy(
                        tsparmcd=parameter.tsparmcd,
                        locator=response.locator,
                        code="CODETABLE_CT_MISMATCH",
                        message=f"Codetable response {response.term_code}/{response.value} "
                        "does not resolve exactly in the selected CT; do not invent a code.",
                    )
                )
    ct_parameters = {t.value for t in terms if t.codelist == "C66738"}
    for code in sorted(ct_parameters ^ {p.tsparmcd for p in snapshot.parameters}):
        discrepancies.append(
            SourceDiscrepancy(
                tsparmcd=code,
                locator="C66738 vs codetable parameter column",
                code="CODETABLE_UNIVERSE_MISMATCH",
                message="Parameter occurs in only one supplied snapshot; review release coverage.",
            )
        )
    return tuple(discrepancies)


class ConformanceRule(Model):
    rule_id: Text
    version: Text
    row: int
    variable: Text
    condition: str | None
    rule: Text
    section: Text
    cited_guidance: Text


def read_ts_conformance(path: Path) -> tuple[ConformanceRule, ...]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook["SDTMIG Conformance Rules v2.0"]
        if next(sheet.iter_rows(values_only=True))[:3] != (
            "Rule ID",
            "SDTMIG Version",
            "Rule Version",
        ):
            raise ValueError("Conformance workbook header mismatch")
        rules = []
        for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if str(row[1]) != "3.4" or row[4] != "TS":
                continue
            rules.append(
                ConformanceRule(
                    rule_id=row[0],
                    version=str(row[2]),
                    row=number,
                    variable=row[5],
                    condition=row[6] or None,
                    rule=row[7],
                    section=str(row[9]),
                    cited_guidance=row[11],
                )
            )
        if not rules or len({r.rule_id for r in rules}) != len(rules):
            raise ValueError("Missing or duplicate TS v3.4 conformance rules")
        return tuple(sorted(rules, key=lambda r: r.rule_id))
    finally:
        workbook.close()
