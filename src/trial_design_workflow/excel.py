"""Excel Output: one draft workbook, three sheets, no intermediate run artifacts."""

import io
import json
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import __version__
from .inputs import StudyInputs
from .ts_registry import CTSnapshot, ParameterCandidate, ParameterRegistry

TS_COLUMNS = ("STUDYID", "DOMAIN", "TSSEQ", "TSPARMCD", "TSPARM", "TSVAL")
REVIEW_COLUMNS = (
    "TSPARMCD",
    "TSPARM",
    "Recommendation",
    "Applicability",
    "Fact_Key",
    "Proposed_Values",
    "Reconciliation",
    "Sources",
    "CT_Result",
    "Review_Status",
    "Reviewer_Action",
    "Reviewer_Notes",
    "Rule_Source",
    "Notes",
)
EVIDENCE_COLUMNS = (
    "Parameter_Fact",
    "TSPARMCD",
    "Field_Key",
    "Proposed_Value",
    "Source_Type",
    "Source_ID",
    "Source_Locator",
    "Evidence",
    "Source_Context",
    "Extraction_Status",
    "CT_Result",
    "Conflict_QC_Message",
    "Review_Status",
)


def _append(sheet, values):
    """Force all source strings to text: no Excel formula injection or silent truncation."""
    row = sheet.max_row + 1 if sheet.max_row > 1 or sheet.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, 1):
        if isinstance(value, str) and (len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value)):
            raise ValueError(
                f"Unrepresentable Excel text at {sheet.title}!{row}:{column}; review input"
            )
        cell = sheet.cell(row, column, value)
        if isinstance(value, str):
            cell.data_type = "s"
            cell.number_format = "@"


def _format(sheet, widths):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "D2" if sheet.title == "TS_Draft" else "C2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        height = 24
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="172B4D")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            width = sheet.column_dimensions[cell.column_letter].width
            text = str(cell.value) if cell.value is not None else ""
            lines = sum(
                max(1, math.ceil(len(line) / max(1, width - 2))) for line in text.split("\n")
            )
            height = max(height, min(409, 15 * lines + 6))
        sheet.row_dimensions[row[0].row].height = height
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def workbook_bytes(
    inputs: StudyInputs,
    registry: ParameterRegistry,
    terminology: CTSnapshot,
    candidates: tuple[ParameterCandidate, ...],
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    draft, review, evidence = (
        workbook.create_sheet(name) for name in ("TS_Draft", "Parameter_Review", "Evidence_QC")
    )
    workbook.properties.title = "Unapproved SDTM TS draft for programmer review"
    workbook.properties.description = "Candidate superset; applicability is not established."
    workbook.properties.creator = "trial-ts"
    for sheet, columns in (
        (draft, TS_COLUMNS),
        (review, REVIEW_COLUMNS),
        (evidence, EVIDENCE_COLUMNS),
    ):
        _append(sheet, columns)
    mapped = set()
    for sequence, candidate in enumerate(candidates, 1):
        parameter, field = candidate.parameter, candidate.field
        _append(
            draft,
            (
                inputs.study_id,
                "TS",
                sequence,
                parameter.TSPARMCD,
                parameter.TSPARM,
                candidate.draft_value,
            ),
        )
        values = [v.proposed_value for v in field.values] if field else []
        sources = (
            sorted({f.source_type for v in field.values for f in v.observations}) if field else []
        )
        _append(
            review,
            (
                parameter.TSPARMCD,
                parameter.TSPARM,
                candidate.recommendation,
                candidate.applicability,
                parameter.expected_fact_key,
                json.dumps(values, ensure_ascii=False),
                field.result if field else "no_fact",
                "; ".join(sources),
                candidate.ct_result,
                "needs_review",
                "CHECK",
                None,
                parameter.rule_source,
                candidate.message + " " + parameter.notes,
            ),
        )
        if field:
            mapped.add(field.field_key)
            for value in field.values:
                for fact in value.observations:
                    _append(
                        evidence,
                        (
                            "parameter_fact",
                            parameter.TSPARMCD,
                            fact.field_key,
                            fact.proposed_value,
                            fact.source_type,
                            fact.source_id,
                            fact.source_locator,
                            fact.evidence_text,
                            fact.context,
                            fact.extraction_status,
                            candidate.ct_result,
                            candidate.message,
                            fact.review_status,
                        ),
                    )
        else:
            _append(
                evidence,
                (
                    "parameter_without_fact",
                    parameter.TSPARMCD,
                    parameter.expected_fact_key,
                    None,
                    None,
                    None,
                    parameter.rule_source,
                    None,
                    parameter.likely_source,
                    "not_extracted",
                    candidate.ct_result,
                    candidate.message,
                    "needs_review",
                ),
            )
    # Preserve all input facts, even if no parameter mapping is implemented yet.
    for field in inputs.facts:
        if field.field_key not in mapped:
            for value in field.values:
                for fact in value.observations:
                    _append(
                        evidence,
                        (
                            "unmapped_fact",
                            None,
                            field.field_key,
                            fact.proposed_value,
                            fact.source_type,
                            fact.source_id,
                            fact.source_locator,
                            fact.evidence_text,
                            fact.context,
                            fact.extraction_status,
                            "not_mapped",
                            f"{field.result}; no automatic TS mapping for this fact",
                            "needs_review",
                        ),
                    )
    for message in (
        "STANDARDS_GAP: SDTMIG 3.4/model/errata and applicability/requiredness rules not loaded. "
        "The CT candidate superset does not establish compliance or study completeness.",
        "DRAFT: All values are proposed, not verified. Empty TSVAL is a review placeholder, "
        "not a null flavor. Delete/edit TS_Draft directly; review actions do not recalculate it.",
        "CT_SCOPE: Only the codelists named in the bundled source metadata are loaded; "
        "unmapped response codelists remain unverified.",
        "CONTEXT: Planned versus actual enrollment and date types are retained separately; "
        "they must not be substituted automatically.",
        *inputs.messages,
    ):
        _append(
            evidence,
            (
                "run_qc",
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                message,
                "needs_review",
            ),
        )
    summary = {
        "input_type": inputs.input_type,
        "nct_id": inputs.nct_id,
        "protocol_sha256": inputs.protocol_sha256,
        "payload_sha256": inputs.payload_sha256,
        "registry_fetched_at": inputs.fetched_at.isoformat() if inputs.fetched_at else None,
        "registry_version": registry.version,
        "registry_sha256": registry.sha256,
        "ct_release": terminology.release,
        "ct_sha256": terminology.sha256,
        "software_version": __version__,
        "study_id": inputs.study_id,
    }
    for key, value in summary.items():
        _append(
            evidence,
            (
                "run_summary",
                None,
                key,
                value,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                "recorded",
            ),
        )
    _format(draft, [27, 10, 11, 15, 46, 88])
    _format(review, [15, 38, 24, 42, 27, 65, 20, 25, 42, 19, 20, 40, 55, 90])
    _format(evidence, [27, 15, 28, 55, 25, 38, 70, 100, 55, 20, 42, 85, 19])
    review.sheet_properties.tabColor = "D97706"
    draft.sheet_properties.tabColor = "243B53"
    validation = DataValidation(type="list", formula1='"KEEP,DELETE,CHECK"')
    validation.errorTitle = "Choose a review action"
    validation.error = "Use KEEP, DELETE, or CHECK. This is a note, not automatic approval."
    validation.showErrorMessage = True
    review.add_data_validation(validation)
    if len(candidates):
        validation.add(f"K2:K{len(candidates) + 1}")
        for row in review.iter_rows(min_row=2, min_col=11, max_col=12):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    target = io.BytesIO()
    workbook.save(target)
    workbook.close()
    content = target.getvalue()
    verify_workbook(content, len(candidates))
    return content


def verify_workbook(content: bytes, candidate_count: int) -> None:
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    try:
        if workbook.sheetnames != ["TS_Draft", "Parameter_Review", "Evidence_QC"]:
            raise ValueError("Workbook sheet validation failed")
        for sheet, columns in (
            ("TS_Draft", TS_COLUMNS),
            ("Parameter_Review", REVIEW_COLUMNS),
            ("Evidence_QC", EVIDENCE_COLUMNS),
        ):
            if tuple(c.value for c in workbook[sheet][1]) != columns:
                raise ValueError(f"Workbook columns invalid in {sheet}")
        rows = list(workbook["TS_Draft"].iter_rows(min_row=2, values_only=True))
        if (
            len(rows) != candidate_count
            or workbook["Parameter_Review"].max_row != candidate_count + 1
        ):
            raise ValueError("Workbook candidate row count validation failed")
        if [r[2] for r in rows] != list(range(1, candidate_count + 1)):
            raise ValueError("Workbook sequence validation failed")
        if len({r[3] for r in rows}) != len(rows) or any(r[1] != "TS" or not r[0] for r in rows):
            raise ValueError("Workbook logical key/structural validation failed")
        if any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row):
            raise ValueError("Unexpected formula in literal draft workbook")
    finally:
        workbook.close()


def save_workbook(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation protects both prior runs and programmer edits.
    with path.open("xb") as stream:
        stream.write(content)
